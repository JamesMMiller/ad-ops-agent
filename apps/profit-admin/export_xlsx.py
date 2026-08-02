"""Build a structured Excel workbook from a warehouse calculate payload."""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


_HEADER_FILL = PatternFill("solid", fgColor="1F2937")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
_TITLE_FONT = Font(bold=True, size=14)
_SECTION_FONT = Font(bold=True, size=12)
_KPI_FILL = PatternFill("solid", fgColor="F3F4F6")
_GOOD_FILL = PatternFill("solid", fgColor="DCFCE7")
_BAD_FILL = PatternFill("solid", fgColor="FEE2E2")
_THIN = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB"),
)


def _autosize(ws, min_width: int = 10, max_width: int = 48) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = 0
        for cell in col:
            if cell.value is None:
                continue
            length = max(length, len(str(cell.value)))
        ws.column_dimensions[letter].width = max(min_width, min(max_width, length + 2))


def _write_header_row(ws, row: int, headers: list[str]) -> None:
    for i, h in enumerate(headers, 1):
        cell = ws.cell(row, i, h)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = _THIN


def _kv_block(ws, start_row: int, pairs: list[tuple[str, Any]], *, title: str | None = None) -> int:
    r = start_row
    if title:
        ws.cell(r, 1, title).font = _SECTION_FONT
        r += 1
    for k, v in pairs:
        ws.cell(r, 1, k).font = Font(bold=True)
        ws.cell(r, 2, v)
        ws.cell(r, 1).border = _THIN
        ws.cell(r, 2).border = _THIN
        r += 1
    return r + 1


def _pct(v: Any) -> str | None:
    if v is None:
        return None
    try:
        return f"{round(float(v) * 100)}%"
    except (TypeError, ValueError):
        return None


def _x(v: Any) -> str | None:
    if v is None:
        return None
    try:
        return f"{float(v):.2f}×"
    except (TypeError, ValueError):
        return None


def build_workbook(payload: dict[str, Any], *, label: str | None = None) -> bytes:
    wb = Workbook()
    inputs = payload.get("inputs") or {}
    unit = payload.get("unit_economics") or {}
    totals = payload.get("totals") or {}
    op = payload.get("overall_pnl") or {}
    past = payload.get("past_performance")
    title = label or inputs.get("label") or "Warehouse plan"

    # ── Summary ──────────────────────────────────────────────
    ws = wb.active
    ws.title = "Summary"
    ws.cell(1, 1, "Our Tech — Profit Admin export").font = _TITLE_FONT
    ws.cell(2, 1, title)
    ws.cell(3, 1, f"Exported {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}")
    ws.cell(4, 1, payload.get("disclaimer") or "")
    ws.merge_cells("A4:F4")
    ws["A4"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[4].height = 48

    kpis = [
        ("Overall P&L (£)", op.get("profit_gbp")),
        ("Margin (inc ads)", _pct(unit.get("margin_after_ads"))),
        ("ROAS", _x(totals.get("roas") or unit.get("roas"))),
        ("ROAS (after COGS)", _x(totals.get("roas_after_cogs") or unit.get("roas_after_cogs"))),
        ("Contrib after ads (£/unit)", unit.get("contribution_after_ads_gbp")),
        ("Landed COGS (£/unit)", unit.get("landed_cogs_gbp")),
        ("POAS", _x(totals.get("poas") or unit.get("poas"))),
        ("Net margin", _pct(totals.get("net_margin"))),
        ("Units sold", totals.get("units_sold")),
        ("Revenue (£)", totals.get("revenue_gbp")),
        ("Ads (£)", totals.get("ads_gbp")),
    ]
    r = 6
    ws.cell(r, 1, "Planner KPIs").font = _SECTION_FONT
    r += 1
    _write_header_row(ws, r, ["Metric", "Value"])
    r += 1
    for name, val in kpis:
        ws.cell(r, 1, name)
        cell = ws.cell(r, 2, val)
        cell.fill = _KPI_FILL
        if name == "Overall P&L (£)" and isinstance(val, (int, float)):
            cell.fill = _GOOD_FILL if val >= 0 else _BAD_FILL
        ws.cell(r, 1).border = _THIN
        cell.border = _THIN
        r += 1

    r += 1
    r = _kv_block(
        ws,
        r,
        [
            ("ROAS", "revenue ÷ ads"),
            ("ROAS (after COGS)", "(revenue − landed COGS) ÷ ads — contribution before ads ÷ ads"),
            ("POAS", "profit ÷ ads (after COGS, fees, storage, ads)"),
            ("Margin (inc ads)", "profit ÷ revenue (or contrib after ads ÷ sell)"),
            ("MER (store desk)", "Shopify revenue ÷ Meta spend — not the same as Meta ROAS"),
            ("Sawtooth on chart", "Restock events — inbound fees + inventory jump"),
        ],
        title="Formulas",
    )

    mode = "Wholesale/dropship (stock tracked)" if inputs.get("track_stock") else "Stock ignored (open demand)"
    r = _kv_block(
        ws,
        r,
        [
            ("Mode", mode),
            ("Warehouse", payload.get("warehouse")),
            ("Stock lane", inputs.get("stock_lane_label")),
            ("Postage in COGS", "Included" if inputs.get("include_postage") else "Excluded"),
            ("Daily ad spend (£)", inputs.get("daily_ad_spend_gbp")),
            ("Ad cost / purchase (£)", inputs.get("ad_cost_per_purchase_gbp")),
            ("Stop ads at stock 0", inputs.get("stop_ads_when_stock_zero")),
            ("Restock", f"{inputs.get('restock_qty')} every {inputs.get('restock_every_days')}d" if inputs.get("restock_enabled") else "Off"),
            ("Source", payload.get("source")),
        ],
        title="Plan context",
    )
    _autosize(ws)

    # ── Inputs ───────────────────────────────────────────────
    ws_in = wb.create_sheet("Inputs")
    pairs = [(k, v) for k, v in sorted(inputs.items()) if not isinstance(v, (dict, list))]
    if isinstance(inputs.get("dest_mix"), dict):
        for k, v in inputs["dest_mix"].items():
            pairs.append((f"dest_mix.{k}", v))
    _write_header_row(ws_in, 1, ["Input", "Value"])
    for i, (k, v) in enumerate(pairs, 2):
        ws_in.cell(i, 1, k)
        ws_in.cell(i, 2, v if not isinstance(v, bool) else ("Yes" if v else "No"))
    _autosize(ws_in)

    # ── Unit economics ───────────────────────────────────────
    ws_u = wb.create_sheet("Unit economics")
    _write_header_row(ws_u, 1, ["Item", "£ / unit", "Notes"])
    unit_rows = [
        ("Sell price", unit.get("sell_price_gbp"), ""),
        ("Product cost", unit.get("product_cost_gbp"), ""),
        ("Inbound (allocated)", unit.get("inbound_alloc_gbp"), "Stock mode"),
        ("Storage (allocated)", unit.get("storage_alloc_gbp"), "Stock mode"),
        ("Outbound fee", unit.get("outbound_gbp"), ""),
        ("Postage in COGS", unit.get("last_mile_gbp"), "0 if postage excluded"),
        ("Landed COGS", unit.get("landed_cogs_gbp"), ""),
        ("Checkout fee", unit.get("checkout_fee_gbp"), ""),
        ("Contribution (ex ads)", unit.get("contribution_gbp"), ""),
        ("Margin (ex ads)", _pct(unit.get("margin")), ""),
        ("Ads / sold unit", unit.get("ads_alloc_gbp"), ""),
        ("Contribution after ads", unit.get("contribution_after_ads_gbp"), ""),
        ("Margin (inc ads)", _pct(unit.get("margin_after_ads")), ""),
        ("ROAS", _x(unit.get("roas")), ""),
        ("ROAS (after COGS)", _x(unit.get("roas_after_cogs")), ""),
        ("POAS", _x(unit.get("poas")), ""),
    ]
    for i, (name, val, note) in enumerate(unit_rows, 2):
        ws_u.cell(i, 1, name)
        ws_u.cell(i, 2, val)
        ws_u.cell(i, 3, note)
    _autosize(ws_u)

    # ── Overall P&L ──────────────────────────────────────────
    ws_p = wb.create_sheet("Overall PnL")
    costs = op.get("costs") or {}
    _write_header_row(ws_p, 1, ["Line", "£"])
    pnl_rows = [
        ("Revenue", op.get("revenue_gbp")),
        ("Product COGS (sold)", costs.get("product_cogs_gbp")),
        ("Inbound", costs.get("inbound_gbp")),
        ("Storage", costs.get("storage_gbp")),
        ("Outbound", costs.get("outbound_gbp")),
        ("Postage", costs.get("postage_gbp")),
        ("Checkout fees", costs.get("checkout_fees_gbp")),
        ("Ads", costs.get("ads_gbp")),
        ("Total costs", op.get("total_costs_gbp")),
        ("Profit / Loss", op.get("profit_gbp")),
        ("Inventory residual (memo)", op.get("inventory_residual_gbp")),
        ("ROAS", _x(op.get("roas"))),
        ("ROAS (after COGS)", _x(op.get("roas_after_cogs"))),
        ("POAS", _x(op.get("poas"))),
        ("Net margin", _pct(op.get("net_margin"))),
    ]
    for i, (name, val) in enumerate(pnl_rows, 2):
        ws_p.cell(i, 1, name).font = Font(bold=name in ("Revenue", "Total costs", "Profit / Loss"))
        cell = ws_p.cell(i, 2, val)
        if name == "Profit / Loss" and isinstance(val, (int, float)):
            cell.fill = _GOOD_FILL if val >= 0 else _BAD_FILL
    ws_p.cell(len(pnl_rows) + 3, 1, op.get("note") or "")
    ws_p.merge_cells(start_row=len(pnl_rows) + 3, start_column=1, end_row=len(pnl_rows) + 3, end_column=2)
    _autosize(ws_p)

    # ── Milestones ───────────────────────────────────────────
    ws_m = wb.create_sheet("Milestones")
    miles = payload.get("milestones") or []
    headers = [
        "Label",
        "Day",
        "Cum sold",
        "Remaining",
        "Cum storage £",
        "Cum revenue £",
        "Cum contrib £",
        "Cum ads £",
        "Cum P&L £",
    ]
    _write_header_row(ws_m, 1, headers)
    for i, m in enumerate(miles, 2):
        ws_m.cell(i, 1, m.get("label"))
        ws_m.cell(i, 2, m.get("day"))
        ws_m.cell(i, 3, m.get("cum_sold"))
        ws_m.cell(i, 4, m.get("remaining"))
        ws_m.cell(i, 5, m.get("cum_storage_gbp"))
        ws_m.cell(i, 6, m.get("cum_revenue_gbp"))
        ws_m.cell(i, 7, m.get("cum_contrib_gbp"))
        ws_m.cell(i, 8, m.get("cum_ads_gbp"))
        ws_m.cell(i, 9, m.get("cum_pnl_gbp"))
    _autosize(ws_m)

    # ── Daily ────────────────────────────────────────────────
    ws_d = wb.create_sheet("Daily")
    days = payload.get("days") or []
    d_headers = [
        "Day",
        "Sold",
        "Restocked",
        "Remaining",
        "Cum sold",
        "Revenue £",
        "Day COGS £",
        "Ads £",
        "Day contrib £",
        "Day P&L £",
        "Cum revenue £",
        "Cum ads £",
        "Cum contrib £",
        "Cum P&L £",
        "Cum storage £",
    ]
    _write_header_row(ws_d, 1, d_headers)
    for i, d in enumerate(days, 2):
        ws_d.cell(i, 1, d.get("day"))
        ws_d.cell(i, 2, d.get("sold"))
        ws_d.cell(i, 3, d.get("restocked") or 0)
        ws_d.cell(i, 4, d.get("remaining"))
        ws_d.cell(i, 5, d.get("cum_sold"))
        ws_d.cell(i, 6, d.get("revenue_gbp"))
        ws_d.cell(i, 7, d.get("day_cogs_gbp"))
        ws_d.cell(i, 8, d.get("ads_gbp"))
        ws_d.cell(i, 9, d.get("day_contrib_gbp"))
        ws_d.cell(i, 10, d.get("day_pnl_gbp"))
        ws_d.cell(i, 11, d.get("cum_revenue_gbp"))
        ws_d.cell(i, 12, d.get("cum_ads_gbp"))
        ws_d.cell(i, 13, d.get("cum_contrib_gbp"))
        ws_d.cell(i, 14, d.get("cum_pnl_gbp"))
        ws_d.cell(i, 15, d.get("cum_storage_gbp"))
    ws_d.freeze_panes = "A2"
    _autosize(ws_d, max_width=16)

    # ── Restock / inbound one-time ───────────────────────────
    ws_r = wb.create_sheet("Inbound & restock")
    ot = payload.get("one_time_usd") or {}
    r = _kv_block(
        ws_r,
        1,
        [
            ("Inspection $", ot.get("inspection")),
            ("Unload $", ot.get("unload")),
            ("Inbound units $", ot.get("inbound_units")),
            ("Labels $", ot.get("labels")),
            ("Stock shipping $", ot.get("stock_shipping")),
            ("Initial total $", ot.get("total")),
            ("Initial total £", payload.get("one_time_gbp")),
            ("Note", ot.get("note")),
        ],
        title="Initial inbound",
    )
    restock = payload.get("restock") or {}
    r = _kv_block(
        ws_r,
        r,
        [
            ("Enabled", restock.get("enabled")),
            ("Qty", restock.get("qty")),
            ("Every days", restock.get("every_days")),
            ("Events", restock.get("events")),
            ("Units added", restock.get("units_added")),
            ("Days", ", ".join(str(d) for d in (restock.get("days") or []))),
        ],
        title="Restock",
    )
    transit = payload.get("transit") or {}
    lm = transit.get("last_mile") or {}
    _kv_block(
        ws_r,
        r,
        [
            ("Stock CN → warehouse", transit.get("stock_to_warehouse")),
            ("Last mile UK", lm.get("UK")),
            ("Last mile EU", lm.get("EU")),
            ("Last mile US", lm.get("US")),
            ("Last mile CAN", lm.get("CAN")),
        ],
        title="Transit",
    )
    _autosize(ws_r)

    # ── Past performance ─────────────────────────────────────
    if past:
        ws_past = wb.create_sheet("Past performance")
        summary = past.get("summary") or {}
        u_p = summary.get("unit_economics") or {}
        op_p = summary.get("overall_pnl") or {}
        t_p = summary.get("totals") or {}
        shop = past.get("shopify") or {}
        meta = past.get("meta") or {}

        ws_past.cell(1, 1, "Past performance (same KPIs as planner)").font = _TITLE_FONT
        ws_past.cell(2, 1, summary.get("label") or shop.get("label") or "")
        ws_past.cell(3, 1, summary.get("source_note") or "")
        ws_past.cell(4, 1, f"Date preset: {past.get('date_preset')}")

        past_kpis = [
            ("Overall P&L (£)", op_p.get("profit_gbp")),
            ("Margin (inc ads)", _pct(u_p.get("margin_after_ads"))),
            ("ROAS", _x(t_p.get("roas") or u_p.get("roas"))),
            ("ROAS (after COGS)", _x(t_p.get("roas_after_cogs") or u_p.get("roas_after_cogs"))),
            ("Contrib after ads (£/unit)", u_p.get("contribution_after_ads_gbp")),
            ("Landed COGS (£/unit)", u_p.get("landed_cogs_gbp")),
            ("POAS", _x(t_p.get("poas"))),
            ("Shopify units", shop.get("units")),
            ("Shopify revenue (£)", shop.get("revenue")),
            ("Shopify COGS (£)", shop.get("cogs")),
            ("Shopify fees (£)", shop.get("fees")),
            ("Meta spend (£)", meta.get("spend")),
            ("Meta purchases", meta.get("purchases")),
            ("Meta CPA (£)", meta.get("cpa")),
            ("Meta attr ROAS", _x(meta.get("roas"))),
        ]
        _write_header_row(ws_past, 6, ["Metric", "Value"])
        for i, (name, val) in enumerate(past_kpis, 7):
            ws_past.cell(i, 1, name)
            cell = ws_past.cell(i, 2, val)
            if name.startswith("Overall P&L") and isinstance(val, (int, float)):
                cell.fill = _GOOD_FILL if val >= 0 else _BAD_FILL

        # Meta objects
        start = 7 + len(past_kpis) + 2
        ws_past.cell(start, 1, "Selected Meta objects").font = _SECTION_FONT
        _write_header_row(
            ws_past,
            start + 1,
            ["Level", "Name", "Spend £", "Purchases", "CPA £", "ROAS"],
        )
        for i, obj in enumerate(meta.get("by_object") or [], start + 2):
            ws_past.cell(i, 1, obj.get("level"))
            ws_past.cell(i, 2, obj.get("name"))
            ws_past.cell(i, 3, obj.get("spend"))
            ws_past.cell(i, 4, obj.get("purchases"))
            ws_past.cell(i, 5, obj.get("cpa"))
            ws_past.cell(i, 6, _x(obj.get("roas")))

        # Shopify lines sheet
        ws_lines = wb.create_sheet("Past Shopify lines")
        _write_header_row(
            ws_lines,
            1,
            ["Date", "Order", "Product", "SKU", "Variant", "Qty", "Revenue £", "COGS £", "Fees £", "Contrib £"],
        )
        for i, line in enumerate(shop.get("lines") or [], 2):
            ws_lines.cell(i, 1, line.get("date"))
            ws_lines.cell(i, 2, line.get("order"))
            ws_lines.cell(i, 3, line.get("product"))
            ws_lines.cell(i, 4, line.get("sku"))
            ws_lines.cell(i, 5, line.get("variant"))
            ws_lines.cell(i, 6, line.get("qty"))
            ws_lines.cell(i, 7, line.get("revenue"))
            ws_lines.cell(i, 8, line.get("cogs"))
            ws_lines.cell(i, 9, line.get("fees"))
            ws_lines.cell(i, 10, line.get("contrib"))
        _autosize(ws_lines)
        _autosize(ws_past)

        if past.get("warnings"):
            ws_w = wb.create_sheet("Warnings")
            ws_w.cell(1, 1, "Warnings").font = _SECTION_FONT
            for i, w in enumerate(past["warnings"], 2):
                ws_w.cell(i, 1, w)
            for i, w in enumerate(payload.get("warnings") or [], len(past["warnings"]) + 3):
                ws_w.cell(i, 1, w)
            _autosize(ws_w)
    elif payload.get("warnings"):
        ws_w = wb.create_sheet("Warnings")
        for i, w in enumerate(payload["warnings"], 1):
            ws_w.cell(i, 1, w)
        _autosize(ws_w)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
